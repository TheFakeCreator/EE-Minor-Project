import os
from pptx import Presentation
from pptx.util import Inches, Pt

def add_bullet_slide(prs, title_text, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = title_text
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        # Check if it's a sub-bullet (indent)
        if bullet.startswith('  - '):
            p.text = bullet[4:]
            p.level = 1
        elif bullet.startswith('- '):
            p.text = bullet[2:]
            p.level = 0
        else:
            p.text = bullet
            p.level = 0
            
    return slide

def add_image_slide(prs, title_text, image_path, description_bullets=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title only layout
    title = slide.shapes.title
    title.text = title_text
    
    if os.path.exists(image_path):
        # Center the image roughly
        pic = slide.shapes.add_picture(image_path, Inches(1), Inches(1.5), width=Inches(8))
    else:
        print(f"Warning: Image not found {image_path}")
        
    if description_bullets:
        txBox = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(8), Inches(1.5))
        tf = txBox.text_frame
        for i, bullet in enumerate(description_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(14)
            
    return slide

def main():
    os.makedirs('reports/mid-sem-ppt', exist_ok=True)
    prs = Presentation()
    
    # 1. Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Grid Stability Prediction"
    subtitle.text = "Mid-Semester Progress Report (EE Minor Project)\n\nPrepared by:\nSanskar Gupta (Roll No: 23117091)\n[Placeholder Teammate]"
    
    # 2. Problem Statement
    add_bullet_slide(prs, "Problem Statement", [
        "- Power grids must maintain stability after disturbances (faults, outages).",
        "- Traditional time-domain simulations are highly accurate but computationally expensive.",
        "- Goal: Predict whether the system will remain stable or go unstable using ML models, drastically reducing assessment time.",
        "- Central Research Question: Can a model generalize to unseen fault conditions?"
    ])
    
    # 3. Methodology & Technology Stack
    add_bullet_slide(prs, "Methodology & Technology Stack", [
        "- Test System: IEEE 39-bus New England system (Standard benchmark).",
        "- Simulation Engine: ANDES (Open-source, Python-based power system simulator).",
        "- ML Pipeline: Scikit-Learn + XGBoost for baseline modeling.",
        "- Pipeline Architecture:",
        "  - Config Generator \u2192 ANDES Simulation Harness \u2192 Feature Extraction \u2192 Stability Labeling"
    ])
    
    # 4. Literature Highlights
    add_bullet_slide(prs, "Literature Highlights", [
        "- Recent work frames Transient Stability Assessment (TSA) as a binary classification problem based on post-fault trajectories.",
        "- Deep Learning (CNNs/RNNs) extracts spatio-temporal features, while Ensemble methods (XGBoost) are popular for tabular feature robustness.",
        "- Standard Stability Criteria: The maximum relative rotor angle deviation between generators exceeding 180 degrees indicates instability."
    ])
    
    # 5. Scenario Space Definition
    add_bullet_slide(prs, "Scenario Space Definition", [
        "- To evaluate our generalization research question, we programmatically generated a diverse dataset.",
        "- Load Scaling: 80% to 120% of nominal.",
        "- Fault Buses: Selected diverse locations (Buses 2, 16, 26, 29, 39).",
        "- Clearing Times: 0.05s to 0.30s.",
        "- Contingencies: Selective N-1 transmission line outages."
    ])
    
    # 6. Dataset Generation Results
    add_bullet_slide(prs, "Dataset Generation Results", [
        "- 450 scenarios simulated in parallel using multi-processing.",
        "- Runtime: ~5 minutes for full dataset generation.",
        "- Class Balance:",
        "  - 308 Stable cases (1)",
        "  - 142 Unstable cases (0)",
        "- Achieved a realistic 2:1 imbalance, providing an excellent foundation for ML training."
    ])
    
    # 7. EDA: Stability vs Clearing Time
    add_image_slide(prs, "EDA: Stability vs Clearing Time", 
                    'reports/figures/stability_vs_clearing_time.png',
                    ["Physical intuition validated: Longer fault clearing times severely increase the probability of system instability."])
    
    # 8. EDA: Feature Distributions
    add_image_slide(prs, "EDA: Feature Distributions", 
                    'reports/figures/feature_distributions.png',
                    ["Unstable scenarios correspond heavily with deeper minimum system voltages during the transient window."])
                    
    # 9. Roadmap & Next Steps
    add_bullet_slide(prs, "Roadmap & Next Steps", [
        "- Minor Project (This Semester):",
        "  - Train baseline models (XGBoost, Random Forest).",
        "  - Execute the core generalization experiments on held-out conditions.",
        "- Major Project (Next Semester):",
        "  - Predict Critical Clearing Time (CCT) with regression.",
        "  - Apply SHAP explainability.",
        "  - Build the Streamlit interactive demo."
    ])
    
    output_path = 'reports/mid-sem-ppt/Grid_Stability_Prediction_MidSem.pptx'
    prs.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == '__main__':
    main()
